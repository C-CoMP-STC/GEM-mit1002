import os

import cobra
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
from gem_utilities import biomass

# Define paths relative to the script or project root
# It's better practice to define a project root
PROJECT_ROOT = os.path.dirname(os.path.dirname(__file__))

import sys

sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Emu, Pt  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

from tools.media import MEDIA  # noqa: E402
from tools.phenotypes import evaluate_phenotypes  # noqa: E402
from tools.plot_styles import summer_colors  # noqa: E402

DATA_DIR = os.path.join(PROJECT_ROOT, "data")
RESULTS_DIR = os.path.join(PROJECT_ROOT, "scripts", "results")

# Load the media definitions
media_definitions = MEDIA

# Define a dictionary of human-friendly versions of the media names
# The key is the name in the media_definitions dictionary
# The value is the human-friendly name to use the table
media_names = {
    "l1": "L1",
    "mbm": "Minimal Basal Medium (Moran Lab)",
    "promm_no_c": "ProMM",
    "marine_broth_wo_yeast_and_peptone": "Marine Broth",
    "marine_broth_wo_yeast_and_peptone_no_n": "Marine Broth (No Nitrogen)",
    "swm": "Seawater Medium",
}

# Human-readable labels for the verdicts tools/phenotypes.py assigns. The
# vocabulary is deliberately the same one figure 2 plots, so the table and the
# figure cannot drift apart. "Positive" means the model predicted growth.
RESULT_LABELS = {
    "true_positive": "True Positive",
    "true_negative": "True Negative",
    "false_positive": "False Positive",
    "false_negative": "False Negative",
    "unsure": "Unsure",
    "invalid_solve": "Invalid Solve",
    "excluded": "Excluded",
}

# Row order: everything that is scored first, then the rows that cannot be
# interpreted, then the rows excluded for a stated reason. Within a group the
# media/metabolite ordering is kept, so the scored block still reads
# alphabetically.
RESULT_ORDER = {
    "True Positive": 0,
    "True Negative": 0,
    "False Positive": 0,
    "False Negative": 0,
    "Unsure": 1,
    "Invalid Solve": 1,
    "Excluded": 2,
}

# Rows the reader should stop on get a fill; the rest do not. See
# color_result_rows for why true positives and true negatives are left white.
RESULT_FILLS = {
    "mismatch": ("False Positive", "False Negative"),
    "unscored": ("Unsure", "Invalid Solve", "Excluded"),
}
# Light gray for the unscored block, and a tint of the manuscript palette's
# pink for the disagreements, so the table matches the figures.
UNSCORED_FILL = "FFE0E0E0"


def _lighten(hex_color: str, amount: float) -> str:
    """Blend ``#RRGGBB`` that far toward white, as an openpyxl aRGB string.

    The palette colours are chosen to be read as a line or a patch, and are too
    saturated to put black text on top of. Deriving the fill from the palette
    rather than hard-coding a second pink means a palette change carries here.
    """
    channels = [int(hex_color.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)]
    blended = [round(c + amount * (255 - c)) for c in channels]
    return "FF" + "".join(f"{c:02X}" for c in blended)


MISMATCH_FILL = _lighten(summer_colors["pink"], 0.55)

# Eight columns of 61 rows does not fit a portrait page at a readable size, so
# the Word version is landscape and set small. Word tables are not typeset by
# ASM at a fixed size, so this only has to be legible, not exact.
DOCX_FONT_SIZE = Pt(8)

# Column widths are shared out in proportion to how much text each column
# actually holds. The clamps stop the two long free-text columns from taking
# the whole page and stop the narrow ones from collapsing, both in characters.
DOCX_MIN_COL_CHARS = 8
DOCX_MAX_COL_CHARS = 30
DOCX_HEADER_PADDING = 4


def generate_growth_phenotype_report(model: cobra.Model):
    # Evaluate every condition with the shared scorer in tools/phenotypes.py,
    # the same one used by test/test_growth.py and by the figure 2 pipeline in
    # curation_process/run_tests_on_prs.py.
    #
    # This function used to carry its own copy of the simulation loop, which
    # disagreed with that scorer in three ways:
    #   * it added a row's metabolites to the medium only when *every* one of
    #     them had an exchange reaction, so "Methionine, Pyruvate" was simulated
    #     on a carbon-free medium rather than on the pyruvate;
    #   * it never checked the solution status, and the model's forced ATP
    #     maintenance bound makes a medium with no usable carbon infeasible
    #     rather than zero-growth, so those rows reported whatever the solver
    #     had left in the primal -- the negative "growth rates" in the table;
    #   * it therefore could not distinguish an infeasible solve from a genuine
    #     prediction of no growth.
    # See the module docstring of tools/phenotypes.py for the full rationale.
    growth_phenotypes = evaluate_phenotypes(model)

    # Re-expose the two columns under the names the heatmap below expects.
    # An unevaluable solve has predicted=None; render it as "Unsure" (gray)
    # rather than silently as no growth.
    growth_phenotypes["all_ex_rxn_present"] = [
        "No" if missing else "Yes" for missing in growth_phenotypes["missing_exchanges"]
    ]
    growth_phenotypes["pred_growth"] = growth_phenotypes["predicted"].fillna("Unsure")

    # Beautify and save the table
    beautify_table(growth_phenotypes)

    # Plot a categorical heatmap of the growth phenotypes, where the rows
    # are the metabolites and the columns are the experimental and predicted
    # growth phenotypes. Show growth as blue and no growth as orange, and
    # unsure as gray
    # First, make a new dataframe with the metabolites as the rows and the
    # experimental and predicted growth phenotypes as the columns
    # Combine the values of "minimal_media" and "c_source" into one column
    growth_phenotypes["c_source"] = (
        growth_phenotypes["minimal_media"] + " " + growth_phenotypes["c_source"]
    )
    # And set it as the index
    growth_phenotypes = growth_phenotypes.set_index("c_source")

    # Make a dictionary for the phenotypes to numbers
    value_to_int = {"Unsure": 0, "No": 1, "Yes": 2}
    n = len(value_to_int)

    # Create an annotation data frame for the text labels on the heatmap
    annotation_key = {"No": "No Exchange", "Yes": ""}
    annot_df = (
        growth_phenotypes["all_ex_rxn_present"].replace(annotation_key).to_frame()
    )
    annot_df.rename(columns={"all_ex_rxn_present": "FBA"}, inplace=True)
    annot_df["Experimental"] = ""
    # Sort the columns to match the order of the heatmap
    annot_df = annot_df[["Experimental", "FBA"]]

    # Subset the other columns, to have just the growth and predicted growth
    growth_phenotypes = growth_phenotypes[["growth", "pred_growth"]]

    # Rename the columns and the index to be longer/more descriptive
    growth_phenotypes.index.name = "Media/Carbon Source"
    growth_phenotypes = growth_phenotypes.rename(
        columns={
            "growth": "Experimental",
            "pred_growth": "FBA",
        }
    )

    # Make a colormap of specified colors (in numerical order for the phenotypes)
    # cmap = ['gray', '#F18F01', '#399E5A'] # Gray, orange, green
    cmap = ["#5E5E5E", "#FF7D0A", "#024064"]  # C-CoMP gray, orange, and dark blue

    # Dynamically set the figure height based on the number of rows
    fig_height = max(10, len(growth_phenotypes) * 0.4)  # 0.4 inches per row
    # Plot the heatmap
    # Use constrained_layout to prevent cutting off y-axis/colorbar labels
    fig, ax = plt.subplots(
        figsize=(8, fig_height),
        constrained_layout=True,
    )
    sns.heatmap(
        growth_phenotypes.replace(value_to_int),
        cmap=cmap,
        linewidths=4,
        linecolor="white",
        annot=annot_df,
        fmt="",
        annot_kws={"fontsize": 8},  # Smaller font size for annotation
        ax=ax,
    )

    # Modify colorbar:
    colorbar = ax.collections[0].colorbar
    r = colorbar.vmax - colorbar.vmin
    colorbar.set_ticks([colorbar.vmin + r / n * (0.5 + i) for i in range(n)])
    colorbar.set_ticklabels(list(value_to_int.keys()))

    # Move the x-axis labels to the top
    plt.tick_params(
        axis="both",
        which="major",
        labelsize=10,
        labelbottom=False,
        bottom=False,
        top=True,
        labeltop=True,
    )

    # Make sure that every y-tick is shown
    ax.set_yticks([i + 0.5 for i in range(len(growth_phenotypes))])
    ax.set_yticklabels(growth_phenotypes.index, rotation=0)

    # Save the figure
    plt.savefig(os.path.join(RESULTS_DIR, "exp_vs_pred_growth_phenotypes.png"))


def beautify_table(exp_pred_table: pd.DataFrame):
    # Get all of the unique minimal media in the table
    unique_minimal_media = exp_pred_table["minimal_media"].unique()
    # Make a dictionary where the keys are the minimal media and the values are the nitrogen-containing compounds in that media, separated by commas
    n_dict = {}
    # For each unique minimal media, look up the nitrogen-containing compounds in that media and add a column to the table with that information
    for minimal_media in unique_minimal_media:
        # Get the media definition for that minimal media
        media_def = media_definitions[minimal_media]
        # Get the nitrogen-containing compounds in that media definition
        n_containing_compounds = []
        for ex_rxn in media_def:
            # Get the metabolite ID from the reaction ID
            # (remove the "EX_" prefix)
            met_id = ex_rxn[3:]  # Remove "EX_"
            # Try to get the metabolite object from the model
            try:
                met = model.metabolites.get_by_id(met_id)
            except KeyError:
                continue
            # Check if the metabolite contains nitrogen in its formula
            if "N" in met.elements:
                # Ignore vitamins (thiamin and vitamin B12)
                if met.id in ["cpd00305_e0", "cpd03424_e0"]:
                    continue
                # If it does, add the name of the metabolite to the list of nitrogen-containing compounds
                # Remove the " [e0]" from the end of the name, if it exists
                if met.name.endswith(" [e0]"):
                    met_name = met.name[:-5]
                else:
                    met_name = met.name
                n_containing_compounds.append(met_name)
        # Store the list of nitrogen-containing compounds for this minimal media
        # as a string separated by commas
        n_dict[minimal_media] = ", ".join(n_containing_compounds)
    # Add a column to the table with the nitrogen-containing compounds, separated by commas
    exp_pred_table["Medium N Source(s)"] = exp_pred_table["minimal_media"].map(n_dict)

    # Replacing the media names in the minimal_media column with the human-friendly versions
    exp_pred_table["minimal_media"] = exp_pred_table["minimal_media"].map(media_names)

    # Round growth rates below (absolute value) 1e-3 to 0
    # To avoid negative 0s
    exp_pred_table["fba_growth_rate"] = exp_pred_table["fba_growth_rate"].apply(
        lambda x: 0 if abs(x) < 1e-3 else x
    )

    # Round the FBA predicted growth rates to 3 decimal places for easier readability
    exp_pred_table["fba_growth_rate"] = exp_pred_table["fba_growth_rate"].round(3)

    # Spell out the scorer's verdict for each row.
    exp_pred_table["result"] = exp_pred_table["category"].map(RESULT_LABELS)

    # Subset the columns we want
    exp_pred_table = exp_pred_table[
        [
            "minimal_media",
            "Medium N Source(s)",
            "c_source",
            "growth",
            "reference",
            "exclude_reason",
            "fba_growth_rate",
            "result",
        ]
    ]
    # Sort the scored rows to the top and the excluded rows to the bottom, then
    # by minimal media and carbon source within each of those groups.
    exp_pred_table = (
        exp_pred_table.assign(
            _group=exp_pred_table["result"].map(RESULT_ORDER),
        )
        .sort_values(by=["_group", "minimal_media", "c_source"])
        .drop(columns="_group")
    )
    # And rename the columns to be more descriptive
    exp_pred_table = exp_pred_table.rename(
        columns={
            "minimal_media": "Minimal Media",
            "c_source": "Added Metabolite(s)",
            "growth": "Experimental Growth",
            "reference": "Reference",
            "exclude_reason": "Reason for Exclusion",
            "fba_growth_rate": "FBA Predicted Growth Rate",
            "result": "Result",
        }
    )
    # Save as a tsv
    exp_pred_table.to_csv(
        os.path.join(RESULTS_DIR, "known_growth_phenotypes_w_pred.tsv"),
        index=False,
        sep="\t",
    )
    # Save as excel, then shade the rows by their verdict
    xlsx_path = os.path.join(RESULTS_DIR, "known_growth_phenotypes_w_pred.xlsx")
    exp_pred_table.to_excel(xlsx_path, index=False)
    color_result_rows(xlsx_path)

    # And as Word, which is the format ASM wants for a main-text table
    save_as_docx(
        exp_pred_table,
        os.path.join(RESULTS_DIR, "known_growth_phenotypes_w_pred.docx"),
    )


def color_result_rows(path: str):
    """Shade each row of the saved workbook according to its Result value.

    Only the rows a reader should stop on are filled: the disagreements, and
    the rows that are not scored at all. True positives and true negatives are
    the large majority of the table, so filling those too would leave the
    shading carrying no information.

    The colour is never the only signal -- the Result column says the same
    thing in text, so the table survives being printed in black and white, or
    being stripped of formatting by a journal's production process.

    Read back from the sheet rather than from the dataframe so the fill cannot
    drift out of step with the row it is describing.
    """
    workbook = load_workbook(path)
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    result_column = header.index("Result") + 1

    fills = {
        label: PatternFill("solid", start_color=color, end_color=color)
        for group, color in (
            ("mismatch", MISMATCH_FILL),
            ("unscored", UNSCORED_FILL),
        )
        for label in RESULT_FILLS[group]
    }

    for row in sheet.iter_rows(min_row=2, max_col=len(header)):
        fill = fills.get(sheet.cell(row=row[0].row, column=result_column).value)
        if fill is None:
            continue
        for cell in row:
            cell.fill = fill

    workbook.save(path)


def _shade_cell(cell, fill: str):
    """Give a table cell a solid background.

    python-docx has no API for cell shading, so the ``w:shd`` element goes in
    by hand. ``w:val="clear"`` means a plain background; the pattern values
    render as black in some viewers.
    """
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_header_row(row):
    """Mark a table row as a header so Word repeats it on every page."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _column_widths(exp_pred_table: pd.DataFrame, total_width):
    """Share ``total_width`` across the columns in proportion to their content.

    Word's own autofit sizes a column to its widest cell, which here gives most
    of the page to the two long free-text columns and wraps
    "Minimal Basal Medium (Moran Lab)" over three lines in every row of the
    medium. Measuring the content and clamping it keeps the table one page
    shorter. Headers are measured by their longest single word, since a header
    is expected to wrap.
    """
    weights = []
    for column_name in exp_pred_table.columns:
        longest_value = max(
            (len(str(value)) for value in exp_pred_table[column_name] if pd.notna(value)),
            default=0,
        )
        # Headers are bold, so they run wider than a character count suggests;
        # without the padding "Experimental" breaks mid-word in its own column.
        longest_header_word = (
            max(len(word) for word in str(column_name).split()) + DOCX_HEADER_PADDING
        )
        weights.append(
            min(
                max(longest_value, longest_header_word, DOCX_MIN_COL_CHARS),
                DOCX_MAX_COL_CHARS,
            )
        )
    total_weight = sum(weights)
    return [Emu(round(total_width * weight / total_weight)) for weight in weights]


def save_as_docx(exp_pred_table: pd.DataFrame, path: str):
    """Write the table as a Word document.

    ASM wants main-text tables in Microsoft Word format, not Excel -- Excel is
    only accepted for supplemental material. This writes the same dataframe the
    TSV and the workbook come from, so the three cannot disagree.

    No caption is written. The table legend is manuscript text and belongs to
    the author; add it in Word above the table.

    Shading matches the workbook, but production restyles main-text tables, so
    it may not survive typesetting. The Result column says the same thing in
    text, which is what the table actually relies on.
    """
    document = Document()

    # Landscape, US Letter. python-docx does not swap the page dimensions when
    # the orientation changes, so both have to be set.
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (
        section.page_height,
        section.page_width,
    )

    widths = _column_widths(
        exp_pred_table,
        section.page_width - section.left_margin - section.right_margin,
    )

    table = document.add_table(rows=1, cols=len(exp_pred_table.columns))
    table.style = "Table Grid"
    # Word only honours the widths if autofit is off and every cell carries
    # one, so they are set per cell as each row is built as well.
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width

    header_cells = table.rows[0].cells
    for cell, column_name, width in zip(header_cells, exp_pred_table.columns, widths):
        cell.width = width
        run = cell.paragraphs[0].add_run(str(column_name))
        run.bold = True
        run.font.size = DOCX_FONT_SIZE
    _repeat_header_row(table.rows[0])

    # The fill is keyed off the Result column exactly as color_result_rows does,
    # so the Word file and the workbook shade the same rows.
    fills = {
        label: color
        for group, color in (
            ("mismatch", MISMATCH_FILL[2:]),
            ("unscored", UNSCORED_FILL[2:]),
        )
        for label in RESULT_FILLS[group]
    }

    for record in exp_pred_table.to_dict("records"):
        cells = table.add_row().cells
        for cell, column_name, width in zip(cells, exp_pred_table.columns, widths):
            cell.width = width
            value = record[column_name]
            text = "" if pd.isna(value) else str(value)
            cell.paragraphs[0].add_run(text).font.size = DOCX_FONT_SIZE
            fill = fills.get(record["Result"])
            if fill is not None:
                _shade_cell(cell, fill)

    document.save(path)


def generate_biomass_producibility_report(model: cobra.Model):
    # Load the TSV of the growth phenotypes
    growth_phenotypes = pd.read_csv(
        os.path.join(DATA_DIR, "known_growth_phenotypes.tsv"),
        sep="\t",
        converters={"met_id": lambda x: x.split(",")},
    )

    # Filter the growth phenotypes to only include the carbon sources that it can grow on
    growth_phenotypes = growth_phenotypes[growth_phenotypes["growth"] == "Yes"]

    # Run the biomass producibility function on each of the models
    sink_options = [False, True]
    for option in sink_options:
        biomass.check_biomass_producibility(
            model,
            growth_phenotypes,
            media_definitions,
            sinks_for_all=option,
            out_dir=RESULTS_DIR,
        )


if __name__ == "__main__":
    # Ensure the results directory exists
    os.makedirs(RESULTS_DIR, exist_ok=True)

    # Load the model
    model = cobra.io.read_sbml_model(os.path.join(PROJECT_ROOT, "model.xml"))

    # Generate the reports
    generate_growth_phenotype_report(model)
    generate_biomass_producibility_report(model)
